"""Create editable slides and a matching PDF from actual project artifacts."""
import json
from datetime import datetime,timezone
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.utils import simpleSplit
ROOT=Path(__file__).resolve().parents[1]
NAVY='143552';BLUE='26789E';PALE='E8F1F7';GRAY='526575'

def read(path):
    file=ROOT/path
    return json.loads(file.read_text(encoding='utf-8-sig')) if file.exists() else None

def slides():
    run=read('models/layoutlmv3-gpu-v2/training_run.json') or {}
    validation=read('models/layoutlmv3-gpu-v2/validation_results.json')
    test=read('models/layoutlmv3-gpu-v2/test_results.json')
    oldv=read('reports/baseline_full_evaluation/validation_results.json')
    oldt=read('reports/baseline_full_evaluation/test_results.json')
    comparison=read('reports/v2_comparison/comparison.json')
    operating=read('run/preferred_mode.json')
    selected_test=read('reports/recommended_mode_test/evaluation.json')
    suites=read('reports/upgraded_tests.json') or []
    external=read('reports/external_kirana_token/evaluation.json')
    policy=read('run/training_policy.json') or {}
    complete=bool(run.get('completed_optimizer_steps') and validation and test)
    result=[]
    def add(title,lines,notes):result.append({'title':title,'lines':lines,'notes':notes})
    add('Smart Document and Finance Analyzer',[
        'Evidence-based financial document extraction and review',
        'Priyanshu Modak (I080) | Vishesh Dokania (I019) | Parth Zanwar (I074)',
        'Native PDFs + OCR | Layout-aware fields | Financial checks | Review and export'],
        'Introduce the team and explain the problem before opening the application. Generated '+datetime.now(timezone.utc).isoformat())
    add('Problem and intended result',[
        'Financial PDFs and scans require repeated manual entry.',
        'Recover fields and rows while retaining their source evidence.',
        'Check arithmetic, expose conflicts and support human correction.',
        'Export structured JSON and nine-sheet Excel workbooks.'],
        'Do not promise fully autonomous accounting or perfect extraction. The application is a local academic review tool.')
    add('Processing pipeline',[
        '1. Validate upload -> embedded PDF words or scan OCR.',
        '2. LayoutLMv3 / QA candidates + explicit field rules.',
        '3. Native tables or header/column geometry -> financial checks.',
        '4. Conservative page grouping -> saved evidence -> review/export.'],
        'Explain that each PDF page independently selects native text or OCR. Table extraction is geometry-based, not trained TATR.')
    add('Datasets and provenance',[
        'Training corpus: 1,000 fictional Indian GST invoices, split 700/150/150.',
        'Exact word boxes and 29 BIO/O labels for 14 financial fields.',
        'Held-out templates: training 0-5, validation 6, test 7.',
        'Extra corpus: 500 CC0 synthetic Kirana invoices; not used in this run.',
        'CORD sample: six real Indonesian receipts; no full benchmark claim.'],
        'Distinguish synthetic from real data. Kirana JSON lacks verified BIO labels and includes visible rounding differences. Cite the dataset card and audit report.')
    add('Training and checkpoint selection',[
        'Isolated CUDA environment: RTX 3050 4 GB, mixed precision.',
        'Batch 1, 256-token word-aligned windows, encoder checkpointing.',
        'Maximum three epochs; predeclared validation stopping target 0.99.',
        ('Completed: '+str(run.get('completed_optimizer_steps'))+' updates, '+str(run.get('completed_epochs'))+' epochs.') if complete else 'Training/evaluation is still pending; final results are not claimed.',
        'Select by validation F1; evaluate the held-out test separately.'],
        'Report actual completion, not requested epochs. The baseline was warm-started for the GPU run. Never select or tune using test scores.')
    metrics=[]
    if oldv and oldt:metrics.append(f"Matched baseline (150/split): validation {oldv['validation_f1']:.3f}, test {oldt['test_f1']:.3f}.")
    else:metrics.append('Matched baseline evaluation: pending.')
    if validation and test:metrics.append(f"GPU model (150/split): validation {validation['validation_f1']:.3f}, test {test['test_f1']:.3f}.")
    else:metrics.append('Final GPU model validation/test metrics: pending.')
    metrics+=['BIO entity F1 uses supplied words and boxes; it is not OCR accuracy.','Historical 20-example baseline scores use a different sample size.']
    add('Measured entity performance',metrics,'Use only the matched full-split rows for a direct comparison. The old small-subset values are historical, and all these data are synthetic.')
    if (ROOT/'reports/gpu_training_curve.pdf').exists():
        result.append({'title':'Recorded training and validation curve','lines':[],'image':'reports/gpu_training_curve.pdf','notes':'Measured saved logs and captured process updates only. Training loss is logarithmic. A lower training loss does not imply better held-out generalization.'})
    rows=[]
    if comparison:
        for mode,row in comparison['modes'].items():rows.append(f"{mode}: field F1 {row['fields']['f1']:.3f}; accuracy {row['fields']['accuracy']:.3f}; exact row F1 {row['line_item_rows']['f1']:.3f}.")
    else:rows.append('Matched OCR-only / model-only / combined evaluation: pending.')
    rows+=['Same ten reserved synthetic test documents in every mode.','Tables use shared geometry; combined fields may come from rules.']
    add('Actual pipeline comparison',rows,'Show final field-source counts and explain why combined accuracy is not standalone neural-model accuracy. Do not retune on this reserved slice.')
    if operating:
        mode_lines=[f"Recommended mode: {operating['recommended_mode']} (matching validation documents)."]
        for name,row in operating['scores'].items():mode_lines.append(f"{name}: validation final-field F1 {row['f1']:.3f}; accuracy {row['accuracy']:.3f}.")
        if selected_test:mode_lines.append(f"Separate 20-document reserved test: field F1 {selected_test['micro']['f1']:.3f}; accuracy {selected_test['micro']['accuracy']:.3f}.")
        mode_lines.append('The model remains available; higher entity F1 did not beat rules end to end.')
        add('Recommended mode and separate test',mode_lines,'Mode choice used validation values. The initial diagnostic test comparison motivated checking modes; no model weights/rules were altered, and a separate reserved test slice was measured afterward.')
    add('Source evidence and review',[
        'Select a field to highlight its source words on the processed page.',
        'Retain candidate alternatives, confidence and review reasons.',
        'Save corrections with human-review provenance and original values.',
        'Explicit approval exports aligned train-only annotations later.'],
        'Show a highlight in the live app. Corrected text absent from OCR may have no highlight and stays pending for token annotation. Approval is separate from immediate training.')
    add('Live demo: digital PDF and continuation',[
        'demo_digital_invoice.pdf: one row, subtotal 100, GST 18, total 118.',
        'Show native PDF text, source highlight and financial checks.',
        'demo_continued_invoice.pdf: two pages -> one invoice, two rows.',
        'Combined subtotal 300, GST 54, total 354; show page provenance.'],
        'Upload and analyze both files. Explain why page 1 may lack totals while the combined invoice reconciles. Both are synthetic demonstrations.')
    add('Live demo: correction, bank and export',[
        'On a fresh synthetic demo, change total 118 -> 120 to show a mismatch.',
        'Restore 118 from the source, save, and show revalidation.',
        'Bank statement: edit transactions and check consecutive balances.',
        'Download JSON and Excel; inspect provenance and merged sheets.'],
        'Clearly state that the deliberate edit is a demonstration of validation, not a model error. Bank/report rows are heuristic extensions.')
    testlines=[f"{suite['suite']}: {suite['tests']} tests, {'passed' if suite['passed'] else 'FAILED'}." for suite in suites]
    testlines+=['Live HTTP uploads verify model execution, source previews and Excel.','Functional checks cover native/scanned PDFs, tables, grouping and review.','Regression counts are behavior checks, not a real-world accuracy sample.']
    add('Verification evidence',testlines,'Point to reports/test_logs, live_app_verification.json and live_upgrade_verification.json. Endpoint health alone does not prove inference is using the selected checkpoint.')
    ext=['External OCR sample: 3/20 field values correct across five validation images.']
    if external:ext.append(f"Combined external sample: field F1 {external['micro']['f1']:.3f}, accuracy {external['micro']['accuracy']:.3f}.")
    else:ext.append('Combined external validation check: pending.')
    ext+=['Different synthetic layouts expose a domain-shift limitation.','Permission-cleared real invoices and manual annotation are still needed.']
    add('Limitations observed in testing',ext,'Do not hide the weak external result or describe synthetic handwriting as verified handwriting recognition. Publisher rounding/partial schemas also limit interpretation.')
    add('Next research steps',[
        'Annotate varied permission-cleared real invoices by vendor/layout.',
        'Reserve independent real validation/test sets before training.',
        'Evaluate learned table structure and additional OCR/language cases.',
        'Add authentication and deployment work before multi-user use.'],
        'These are explicitly future work. Existing tests and synthetic scores do not establish production readiness or tax-registration verification.')
    add('Demonstration resources and references',[
        'Run dashboard: http://127.0.0.1:8501 | API: http://127.0.0.1:8000/docs',
        'docs/START_HERE.md | docs/PROFESSOR_DEMO.md | PROJECT_HANDBOOK.pdf',
        'Measured results: reports/UPGRADE_RESULTS.md; generated status: docs/STATUS.md.',
        'Sources: pdfplumber, Microsoft LayoutLMv3/TATR, PyTorch, CORD and Kirana.',
        'Model license: CC-BY-NC-SA-4.0; retain attribution and usage conditions.'],
        'Primary source links are in docs/REFERENCES.md and the handbook. Close by distinguishing completed work from remaining evaluation.')
    return result

def textbox(slide,text,x,y,w,h,size=24,color=NAVY,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame;frame.word_wrap=True;frame.margin_left=0;frame.margin_right=0
    frame.text=text
    for paragraph in frame.paragraphs:
        paragraph.font.name='Aptos';paragraph.font.size=Pt(size);paragraph.font.bold=bold;paragraph.font.color.rgb=RGBColor.from_string(color)
    return shape

def main():
    pages=slides();deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    for number,page in enumerate(pages,1):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        band=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,deck.slide_width,Inches(.18));band.fill.solid();band.fill.fore_color.rgb=RGBColor.from_string(BLUE);band.line.fill.background()
        textbox(slide,page['title'],.65,.55,12.05,1,34,bold=True)
        if page.get('image'):
            import pypdfium2 as pdfium
            source=pdfium.PdfDocument(str(ROOT/page['image']))
            image=ROOT/'reports/gpu_training_curve_slide.png'
            source[0].render(scale=2).to_pil().save(image)
            slide.shapes.add_picture(str(image),Inches(.65),Inches(1.9),width=Inches(12))
        for index,line in enumerate(page['lines']):
            textbox(slide,str(index+1).zfill(2),.7,1.85+index*.83,.55,.4,16,BLUE,True)
            textbox(slide,line,1.4,1.78+index*.83,11.1,.78,23)
        textbox(slide,'SMART DOCUMENT + FINANCE ANALYZER | Synthetic scores and functional evidence',.7,7.05,11.3,.25,10,GRAY)
        textbox(slide,str(number),12.3,7.05,.4,.25,10,GRAY)
        slide.notes_slide.notes_text_frame.text=page['notes']
    target=ROOT/'artifacts/PROFESSOR_PRESENTATION.pptx';target.parent.mkdir(exist_ok=True);deck.save(target)
    pdf=canvas.Canvas(str(target.with_suffix('.pdf')),pagesize=(960,540))
    for number,page in enumerate(pages,1):
        pdf.setFillColor(colors.HexColor('#'+BLUE));pdf.rect(0,527,960,13,fill=1,stroke=0)
        pdf.setFillColor(colors.HexColor('#'+NAVY));pdf.setFont('Helvetica-Bold',26)
        for i,line in enumerate(simpleSplit(page['title'],'Helvetica-Bold',26,870)):pdf.drawString(48,475-i*31,line)
        if page.get('image'):
            pdf.drawImage(str(ROOT/'reports/gpu_training_curve_slide.png'),48,100,width=864,height=324,mask='auto')
        y=390
        for index,line in enumerate(page['lines']):
            pdf.setFillColor(colors.HexColor('#'+BLUE));pdf.setFont('Helvetica-Bold',12);pdf.drawString(52,y,str(index+1).zfill(2))
            pdf.setFillColor(colors.HexColor('#'+NAVY));pdf.setFont('Helvetica',17)
            wrapped=simpleSplit(line,'Helvetica',17,795)
            for offset,text in enumerate(wrapped):pdf.drawString(98,y-offset*22,text)
            y-=max(60,len(wrapped)*22+12)
        pdf.setFillColor(colors.HexColor('#'+GRAY));pdf.setFont('Helvetica',8);pdf.drawString(48,23,'SMART DOCUMENT + FINANCE ANALYZER | Synthetic scores and functional evidence');pdf.drawRightString(915,23,str(number));pdf.showPage()
    pdf.save()
    # Validate saved structure, page count and text presence without requiring PowerPoint.
    restored=Presentation(target)
    assert len(restored.slides)==len(pages)
    assert all(any(shape.has_text_frame and shape.text.strip() for shape in slide.shapes) for slide in restored.slides)
    (ROOT/'reports/presentation_manifest.json').write_text(json.dumps({'slides':len(pages),'generated_at':datetime.now(timezone.utc).isoformat(),'source':'Measured JSON artifacts and maintained demonstration notes','pending_fields_explicit':True},indent=2))
    print(f'Built {len(pages)} editable slides and matching PDF',flush=True)

if __name__=='__main__':main()
