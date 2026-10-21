import io
import os
import tempfile
import unittest
from unittest.mock import patch
from PIL import Image
from fastapi.testclient import TestClient
from backend.app import app
from backend.layout_extractor import extract_line_items, extract_fields
from backend.ocr_engine import Token

class PipelineTests(unittest.TestCase):
    def test_table_does_not_include_totals_or_description_numbers(self):
        rows = [['Item','Qty','Rate','Amount'],['A4','Paper','2','100.00','200.00'],['Subtotal','200.00']]
        tokens = [{'text':word,'box':[x*100,y*30,x*100+80,y*30+20]} for y,row in enumerate(rows) for x,word in enumerate(row)]
        items = extract_line_items(tokens)
        self.assertEqual(len(items),1)
        self.assertEqual(items[0]['description'],'A4 Paper')
        self.assertEqual(items[0]['amount'],200)

    def test_offline_mode(self):
        with patch.dict(os.environ,{'EXTRACTION_MODE':'ocr'}):
            fields = extract_fields(Image.new('RGB',(10,10)), 'Seller: Test Ltd\nInvoice No: INV-1\nSubtotal: Rs.100\nGST (18%): Rs.18\nTotal: Rs.118')
        self.assertEqual(fields['invoice_number']['value'],'INV-1')
        self.assertEqual(fields['total_amount']['value'],'Rs.118')

    def test_upload_history_correction_and_exports(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ,{'DATABASE_PATH':directory+'/test.sqlite','EXTRACTION_MODE':'ocr'}):
            image = io.BytesIO()
            Image.new('RGB',(200,200),'white').save(image,format='PNG')
            with TestClient(app) as client:
                self.assertEqual(client.post('/extract',files={'file':('empty.png',b'')}).status_code,400)
                self.assertEqual(client.post('/extract',files={'file':('bad.exe',b'123')}).status_code,400)
                token = Token('Seller',[0,0,50,20],[0,0,250,100],95)
                with patch('backend.app.run_ocr',return_value=[token]),patch('backend.app.preprocess_page',side_effect=lambda page:page):
                    response = client.post('/extract',files={'file':('invoice.png',image.getvalue(),'image/png')})
                self.assertEqual(response.status_code,200,response.text)
                result = response.json()
                self.assertEqual(result['status'],'review_needed')
                doc_id = result['id']
                self.assertEqual(len(client.get('/documents').json()['documents']),1)
                corrected = client.patch('/documents/'+doc_id,json={'page':1,'fields':{'vendor_name':'=HYPERLINK("bad")'}})
                self.assertEqual(corrected.status_code,200)
                self.assertEqual(corrected.json()['pages'][0]['fields']['vendor_name']['source'],'human_review')
                excel = client.get('/documents/'+doc_id+'/excel')
                self.assertTrue(excel.content.startswith(b'PK'))
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(excel.content))
                self.assertEqual(set(book.sheetnames),{'Fields','Line Items','Validation','Provenance','Transactions','Financial Metrics','Spatial Relations','Document Summary','Merged Line Items'})
                for row in book['Fields']:
                    for cell in row:
                        self.assertNotEqual(cell.data_type,'f')
                self.assertEqual(client.get('/documents/not-found').status_code,404)

if __name__=='__main__':
    unittest.main()
