"""The Streamlit deployment uses the same pipeline without HTTP requests."""
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from backend.standalone import request
from backend import app


class StandaloneTests(unittest.TestCase):
    def test_full_workflow_and_isolation_without_network(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {'DATABASE_PATH':directory+'/documents.sqlite3','EXTRACTION_MODE':'ocr'}), patch('requests.request',side_effect=AssertionError('HTTP must not be used')):
            owner={'X-Visitor-Session':'a'*64}
            other={'X-Visitor-Session':'b'*64}
            sample=Path(__file__).resolve().parents[1]/'sample_docs/demo_digital_invoice.pdf'
            self.assertEqual(request('GET','/documents').status_code,401)
            result=request('POST','/extract',headers=owner,files={'file':(sample.name,sample.read_bytes())})
            self.assertEqual(result.status_code,200,result.text)
            identity=result.json()['id']
            self.assertEqual(len(request('GET','/documents',headers=owner).json()['documents']),1)
            for suffix in ('','/excel','/pages/1/preview'):
                self.assertEqual(request('GET','/documents/'+identity+suffix,headers=other).status_code,404)
            self.assertTrue(request('GET','/documents/'+identity+'/pages/1/preview',headers=owner).content.startswith(b'\x89PNG'))
            changed=request('PATCH','/documents/'+identity,headers=owner,json={'page':1,'fields':{'total_amount':'130'}})
            self.assertEqual(changed.status_code,200)
            self.assertEqual(changed.json()['status'],'review_needed')
            restored=request('PATCH','/documents/'+identity,headers=owner,json={'page':1,'fields':{'total_amount':'118'}})
            self.assertEqual(restored.status_code,200)
            workbook=load_workbook(io.BytesIO(request('GET','/documents/'+identity+'/excel',headers=owner).content))
            self.assertEqual(len(workbook.sheetnames),9)
            self.assertEqual(request('GET','/training-reviews',headers=owner).status_code,403)
            self.assertEqual(request('DELETE','/documents',headers=other).json()['deleted'],0)
            self.assertEqual(request('DELETE','/documents',headers=owner).json()['deleted'],1)
            self.assertEqual(request('GET','/documents/'+identity,headers=owner).status_code,404)

    def test_upload_guards_and_invalid_corrections(self):
        headers={'X-Visitor-Session':'a'*64}
        for name,raw,status in [('bad.exe',b'123',400),('empty.pdf',b'',400),('too-large.pdf',b'12345',413)]:
            with patch.object(app,'MAX_BYTES',4):
                self.assertEqual(request('POST','/extract',headers=headers,files={'file':(name,raw)}).status_code,status)
        self.assertEqual(request('PATCH','/documents/'+'a'*32,headers=headers,json={'page':0}).status_code,400)
